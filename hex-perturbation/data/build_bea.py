"""Build `bea_klems_1963_2016.csv` from the BEA-BLS research file (public domain, U.S. government).

Source: Eldridge et al., "Toward a BEA-BLS Integrated Industry-level Production Account for
1947-2016", underlying data
https://www.bea.gov/sites/default/files/2020-11/KLEMS-Research-Tables-Underlying-Data_0.xlsx
(sheet "1963-2016": 63 industries, 54 years).

Per industry and year: nominal values and quantity indices of gross output, intermediate input,
five capital types and two labour types.  Capital and labour are aggregated by a chained Tornqvist
quantity index (weights = average nominal shares of adjacent years), prices = nominal / quantity.

Usage: python build_bea.py <path to KLEMS-Research-Tables-Underlying-Data_0.xlsx>
"""

import csv
import math
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent / "bea_klems_1963_2016.csv"
CAP = [("vkit.", "qkit."), ("vksoft.", "qks."), ("vkRD.", "qkrd."), ("vkart.", "qka."), ("vkoth.", "qko.")]
LAB = [("vlcol.", "qlindexcol_merge."), ("vln.", "qlindexn_merge.")]


def tornqvist(rows, parts):
    """Chained Tornqvist quantity index (first year = 1) and the nominal total, per year."""
    nom = [sum(r[v] for v, _ in parts) for r in rows]
    q = [1.0]
    for a, b, na, nb in zip(rows, rows[1:], nom, nom[1:]):
        g = 0.0
        for v, k in parts:
            if a[k] > 0 and b[k] > 0:
                g += 0.5 * (a[v] / na + b[v] / nb) * math.log(b[k] / a[k])
        q.append(q[-1] * math.exp(g))
    return nom, q


def series(wb, sheet):
    """{(indnum, name): [(year, pK, pL, pII, pGO, GO_Q), ...]} for one sheet."""
    it = wb[sheet].iter_rows(values_only=True)
    next(it)
    head = next(it)
    by_ind = defaultdict(list)
    for row in it:
        if row[0] is None:
            continue
        r = dict(zip(head, row))
        by_ind[(r["indnum"], r["Description"])].append(r)
    out = {}
    for key, rows in by_ind.items():
        rows.sort(key=lambda r: r["yr"])
        vk, qk = tornqvist(rows, CAP)
        vl, ql = tornqvist(rows, LAB)
        out[key] = [(r["yr"], a / b, c / e, r["ii."] / r["iiqi."], r["go."] / r["goqi."], r["goqi."])
                    for r, a, b, c, e in zip(rows, vk, qk, vl, ql)]
    return out


def write(path, data):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["indnum", "industry", "year", "pK", "pL", "pII", "pGO", "GO_Q"])
        for (ind, name), rows in sorted(data.items()):
            for row in rows:
                w.writerow([ind, name, row[0]] + [f"{v:.6g}" for v in row[1:]])
    print(f"wrote {path} ({len(data)} industries)")


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    late = series(wb, "1963-2016")
    write(OUT, late)
    # splice 1947-1963 onto 1963-2016 for industries present in both sheets (same indnum and
    # name): every early series is multiplied by one factor matching the two 1963 values
    early = series(wb, "1947-1963")
    spliced = {}
    for key in sorted(set(early) & set(late)):
        e, l = early[key], late[key]
        assert e[-1][0] == l[0][0] == 1963
        f = [l[0][i] / e[-1][i] for i in range(1, 6)]
        spliced[key] = [(r[0],) + tuple(v * fi for v, fi in zip(r[1:], f)) for r in e[:-1]] + l
    write(OUT.with_name("bea_klems_1947_2016.csv"), spliced)
    write_detail(wb, OUT.with_name("bea_klems_detail_1963_2016.csv"))


def write_detail(wb, path):
    """Disaggregated prices (nominal / quantity index) of the 5 capital and 2 labour types, II, GO,
    and GO_Q, 1963-2016: up to d + 1 = 9 prices.  Types with a zero quantity index get an empty
    price."""
    it = wb["1963-2016"].iter_rows(values_only=True)
    next(it)
    head = next(it)
    parts = CAP + LAB + [("ii.", "iiqi."), ("go.", "goqi.")]
    names = ["pKit", "pKsoft", "pKrd", "pKart", "pKoth", "pLcol", "pLn", "pII", "pGO"]
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["indnum", "industry", "year"] + names + ["GO_Q"])
        rows = sorted((dict(zip(head, row)) for row in it if row[0] is not None),
                      key=lambda r: (r["indnum"], r["yr"]))
        for r in rows:
            prices = [f"{r[v] / r[q]:.6g}" if r[q] and r[q] > 0 and r[v] > 0 else "" for v, q in parts]
            w.writerow([r["indnum"], r["Description"], r["yr"]] + prices + [f"{r['goqi.']:.6g}"])
    print(f"wrote {path}")


if __name__ == "__main__":
    main(sys.argv[1])
